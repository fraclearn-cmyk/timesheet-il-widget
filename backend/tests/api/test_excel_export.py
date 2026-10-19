"""Public XLSX contract for the scoped detailed timesheet."""

from datetime import date, datetime, timedelta
import base64
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import os
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import sessionmaker

from app.core.access_policy import RequestContext
from app.core.database import Base
from app.api.v1.dependencies import APIProblem
from app.models import (
    GroupMember,
    StatusTransition,
    User,
    WidgetGroup,
    WorkSession,
    WorkStatus,
)


PATH = "/api/v1/reports/export-excel"
PERIOD = {"date_from": "2026-09-22", "date_to": "2026-09-22"}


def export(client, **body):
    return client.post(PATH, json={**PERIOD, **body})


def sheet(response):
    assert response.status_code == 200, response.text
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Табель"]
    return book.active


def add_session(
    db,
    *,
    day=date(2026, 9, 22),
    external_id=102,
    account=10,
    start=datetime(2026, 9, 22, 6),
    end=datetime(2026, 9, 22, 7),
    status=WorkStatus.FINISHED,
    work=3600,
    breaks=0,
    late=0,
    transitions=(),
):
    row = WorkSession(
        amocrm_account_id=account,
        amocrm_user_id=external_id,
        user_name="PRIVATE CRM snapshot",
        department="PRIVATE department",
        business_date=day,
        start_time=start,
        end_time=end,
        current_status=status,
        total_work_time=work,
        total_break_time=breaks,
        late_minutes=late,
        late_reason="PRIVATE payload https://example.invalid",
    )
    db.add(row)
    db.flush()
    for state, at in transitions:
        db.add(StatusTransition(work_session_id=row.id, to_status=state, timestamp=at))
    db.flush()
    return row


@pytest.mark.parametrize(
    "field,group,label,value",
    [
        ("shift_lateness", "Начало дня", "Минуты", 7),
        ("shift_late_reason", "Начало дня", "Объяснительная", "'=поезд"),
        ("break_lateness", "Перерыв", "Минуты", 13),
        (
            "break_late_reason",
            "Перерыв",
            "Объяснительная",
            "Перерыв №1 — 8 мин: встреча\nПерерыв №2 — 5 мин: +связь",
        ),
        ("lateness", "Начало дня", "Минуты", 7),
    ],
)
def test_each_lateness_field_is_independent_with_three_header_rows(
    scoped_client, db, field, group, label, value
):
    from app.models.lateness_incident import LatenessIncident, LatenessKind

    session = add_session(db, late=99)
    for kind, sequence, minutes, reason in [
        (LatenessKind.SHIFT_START, 0, 7, "=поезд"),
        (LatenessKind.BREAK_RETURN, 2, 5, "+связь"),
        (LatenessKind.BREAK_RETURN, 1, 8, "встреча"),
    ]:
        db.add(
            LatenessIncident(
                account_id=10,
                amocrm_user_id=102,
                work_session_id=session.id,
                kind=kind,
                sequence=sequence,
                occurred_at=session.start_time,
                late_minutes=minutes,
                reason=reason,
            )
        )
    db.commit()
    ws = sheet(export(scoped_client("admin"), columns=[field]))
    assert [ws.cell(row, 1).value for row in range(1, 5)] == [
        "Опоздания",
        group,
        label,
        value,
    ]
    assert ws.freeze_panes == "A4"
    assert ws.auto_filter.ref == "A3:A4"
    assert ws["A4"].number_format == ("0" if label == "Минуты" else "General")
    if label == "Объяснительная":
        assert ws["A4"].alignment.wrap_text
        assert ws["A4"].data_type == "s"


def test_restarted_shift_reasons_are_visible_in_excel(scoped_client, db):
    from app.models.lateness_incident import LatenessIncident, LatenessKind

    first = add_session(
        db,
        start=datetime(2026, 9, 22, 6),
        end=datetime(2026, 9, 22, 7),
    )
    second = add_session(
        db,
        start=datetime(2026, 9, 22, 8),
        end=datetime(2026, 9, 22, 9),
    )
    for session, minutes, reason in [
        (first, 7, "автобус"),
        (second, 11, "врач"),
    ]:
        db.add(
            LatenessIncident(
                account_id=10,
                amocrm_user_id=102,
                work_session_id=session.id,
                kind=LatenessKind.SHIFT_START,
                sequence=0,
                occurred_at=session.start_time,
                late_minutes=minutes,
                reason=reason,
            )
        )
    db.commit()

    ws = sheet(
        export(
            scoped_client("admin"),
            columns=["shift_lateness", "shift_late_reason"],
        )
    )
    assert ws["A4"].value == 18
    assert ws["B4"].value == "автобус\nврач"
    assert ws["B4"].alignment.wrap_text


def test_default_workbook_groups_headers_and_preserves_minute_units(scoped_client, db):
    add_session(db, late=75)
    db.commit()
    ws = sheet(export(scoped_client("admin")))
    assert {str(merge) for merge in ws.merged_cells.ranges} == {
        "A1:A3",
        "B1:B3",
        "C1:C3",
        "D1:D3",
        "E1:E3",
        "F1:F3",
        "G1:G3",
        "H1:K1",
        "H2:I2",
        "J2:K2",
    }
    assert [ws.cell(3, col).value for col in range(8, 12)] == [
        "Минуты",
        "Объяснительная",
        "Минуты",
        "Объяснительная",
    ]
    assert ws["H4"].value == 75
    assert ws["J4"].value == 0
    assert ws["I4"].value is None
    assert ws.freeze_panes == "A4"
    assert ws.auto_filter.ref == "A3:K4"


def test_partial_groups_only_merge_contiguous_selected_columns(scoped_client, db):
    add_session(db)
    db.commit()
    ws = sheet(
        export(
            scoped_client("admin"),
            columns=[
                "shift_late_reason",
                "break_lateness",
                "employee",
                "shift_lateness",
                "break_late_reason",
            ],
        )
    )
    assert {str(merge) for merge in ws.merged_cells.ranges} == {
        "A1:B1",
        "C1:C3",
        "D1:E1",
    }
    assert [cell.value for cell in ws[2]] == [
        "Начало дня",
        "Перерыв",
        None,
        "Начало дня",
        "Перерыв",
    ]
    assert [cell.value for cell in ws[3]] == [
        "Объяснительная",
        "Минуты",
        None,
        "Минуты",
        "Объяснительная",
    ]
    assert ws["C4"].value == "Employee"


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_structured_explanations_are_formula_safe(scoped_client, db, prefix):
    from app.models.lateness_incident import LatenessIncident, LatenessKind

    session = add_session(db)
    db.add(
        LatenessIncident(
            account_id=10,
            amocrm_user_id=102,
            work_session_id=session.id,
            kind=LatenessKind.SHIFT_START,
            sequence=0,
            occurred_at=session.start_time,
            late_minutes=1,
            reason=f'  {prefix}HYPERLINK("https://evil.invalid")',
        )
    )
    db.commit()
    response = export(scoped_client("admin"), columns=["shift_late_reason"])
    ws = sheet(response)
    # API normalizes surrounding whitespace, then XLSX protects formula prefixes.
    assert ws["A4"].value == f'\'{prefix}HYPERLINK("https://evil.invalid")'
    assert ws["A4"].data_type == "s"
    with ZipFile(BytesIO(response.content)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in xml


def add_late_breaks(db, session, *, count=40, prefix="", character="x"):
    from app.models.lateness_incident import LatenessIncident, LatenessKind

    reason = prefix + character * (999 - len(prefix)) + "\x01"
    for sequence in range(1, count + 1):
        db.add(
            LatenessIncident(
                account_id=session.amocrm_account_id,
                amocrm_user_id=session.amocrm_user_id,
                work_session_id=session.id,
                kind=LatenessKind.BREAK_RETURN,
                sequence=sequence,
                occurred_at=session.start_time,
                late_minutes=sequence,
                reason=reason,
            )
        )
    return reason.replace("\x01", "\ufffd")


@pytest.mark.parametrize(
    "field,kind,sequence",
    [
        ("shift_late_reason", "shift_start", 0),
        ("break_late_reason", "break_return", 1),
    ],
)
def test_xml_illegal_controls_are_replaced_and_supported_whitespace_survives(
    scoped_client,
    db,
    field,
    kind,
    sequence,
):
    from app.models.lateness_incident import LatenessIncident

    db.get(User, 2).name = "=Employee\x01"
    session = add_session(db)
    reason = "=reason\x01with\x00controls\x0b\x1f\ufffe\uffff\tand\nnew\rline"
    expected = "=reason\ufffdwith\ufffdcontrols\ufffd\ufffd\ufffd\ufffd\tand\nnew\rline"
    from app.services.timesheet_excel_service import _safe_text

    assert _safe_text(reason) == "'" + expected
    db.add(
        LatenessIncident(
            account_id=10,
            amocrm_user_id=102,
            work_session_id=session.id,
            kind=kind,
            sequence=sequence,
            occurred_at=session.start_time,
            late_minutes=1,
            reason=reason,
        )
    )
    db.commit()
    ws = sheet(export(scoped_client("admin"), columns=["employee", field]))
    assert ws["A4"].value == "'=Employee\ufffd"
    # XML parsers normalize literal carriage returns to line feeds on reload.
    expected = expected.replace("\r", "\n")
    assert ws["B4"].value == (
        "'" + expected
        if field == "shift_late_reason"
        else f"Перерыв №1 — 1 мин: {expected}"
    )
    assert ws["A4"].data_type == ws["B4"].data_type == "s"


@pytest.mark.parametrize(
    "prefix,character", [("=", "x"), ("+", "x"), ("-", "x"), ("@", "x"), ("=", "😀")]
)
def test_overflow_break_reasons_keep_complete_entries_and_all_incidents_on_auxiliary_sheet(
    scoped_client,
    db,
    prefix,
    character,
):
    db.get(User, 2).name = prefix + "Employee\x01"
    session = add_session(db)
    reason = add_late_breaks(db, session, prefix=prefix, character=character)
    # The auxiliary sheet includes only affected days, although the main report has two.
    other = add_session(
        db,
        day=date(2026, 9, 23),
        start=datetime(2026, 9, 23, 6),
        end=datetime(2026, 9, 23, 7),
    )
    add_late_breaks(db, other, count=1)
    db.commit()
    response = export(
        scoped_client("admin"), date_to="2026-09-23", columns=["break_late_reason"]
    )
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Табель", "Объяснительные"]
    main, details = book.worksheets
    text = main["A5"].value
    lines = text.split("\n")
    assert lines[-1] == "Полный список — на листе «Объяснительные»."
    expected = [
        f"Перерыв №{sequence} — {sequence} мин: {reason}" for sequence in range(1, 41)
    ]
    assert lines[:-1] == expected[: len(lines) - 1]
    assert 1 < len(lines) < 41
    assert len(text.encode("utf-16-le")) // 2 <= 32767
    assert (
        len((text + "\n" + expected[len(lines) - 1]).encode("utf-16-le")) // 2 > 32767
    )
    assert [cell.value for cell in details[1]] == [
        "Сотрудник",
        "Дата",
        "Перерыв №",
        "Минуты",
        "Объяснительная",
    ]
    values = list(details.iter_rows(min_row=2, values_only=True))
    assert values == [
        (
            "'" + prefix + "Employee\ufffd",
            datetime(2026, 9, 22),
            sequence,
            sequence,
            "'" + reason,
        )
        for sequence in range(1, 41)
    ]
    assert values[-1][2] == 40 and values[-1][4] == "'" + reason
    assert details["D41"].number_format == "0"
    assert details["E41"].alignment.wrap_text
    assert details["E41"].data_type == "s"
    assert details.freeze_panes == "A2"
    assert details.auto_filter.ref == "A1:E41"
    assert main["A5"].alignment.wrap_text
    with ZipFile(BytesIO(response.content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = archive.read(name)
                ET.fromstring(xml)
                assert b"<f>" not in xml


@pytest.mark.parametrize(
    "columns,count",
    [
        (["employee"], 40),
        (["break_lateness"], 40),
        (["shift_late_reason"], 40),
        (["break_late_reason"], 1),
    ],
)
def test_auxiliary_sheet_is_absent_when_overflow_is_unselected_or_unnecessary(
    scoped_client,
    db,
    columns,
    count,
):
    add_late_breaks(db, add_session(db), count=count)
    db.commit()
    ws = sheet(export(scoped_client("admin"), columns=columns))
    assert "Объяснительные" not in "".join(
        str(cell.value or "") for cells in ws for cell in cells
    )


@pytest.mark.parametrize(
    "start,end,days",
    [
        ("2026-09-22", "2026-09-22", [date(2026, 9, 22)]),
        ("2026-09-22", "2026-09-28", [date(2026, 9, 22), date(2026, 9, 28)]),
        ("2026-09-01", "2026-09-30", [date(2026, 9, 22), date(2026, 9, 30)]),
        ("2026-06-22", "2026-09-22", [date(2026, 6, 22), date(2026, 9, 22)]),
    ],
)
def test_periods_include_exact_employee_days(scoped_client, db, start, end, days):
    add_session(db)
    if end == "2026-09-28":
        add_session(
            db,
            day=date(2026, 9, 28),
            start=datetime(2026, 9, 28, 6),
            end=datetime(2026, 9, 28, 7),
        )
    if end == "2026-09-30":
        add_session(
            db,
            day=date(2026, 9, 30),
            start=datetime(2026, 9, 30, 6),
            end=datetime(2026, 9, 30, 7),
        )
    if start == "2026-06-22":
        add_session(
            db,
            day=date(2026, 6, 22),
            start=datetime(2026, 6, 22, 6),
            end=datetime(2026, 6, 22, 7),
        )
    db.commit()
    response = export(scoped_client("admin"), date_from=start, date_to=end)
    ws = sheet(response)
    assert [cell.value.date() for cell in ws["B"][3:]] == list(reversed(days))
    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="timesheet_{start}_{end}.xlsx"'
    )


def test_json_export_wraps_the_same_workbook_for_amocrm_transport(scoped_client, db):
    add_session(db)
    db.commit()

    response = scoped_client("admin").post(f"{PATH}-json", json=PERIOD)

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["filename"] == "timesheet_2026-09-22_2026-09-22.xlsx"
    workbook = load_workbook(BytesIO(base64.b64decode(payload["content_base64"])))
    assert workbook.sheetnames == ["Табель"]
    assert workbook.active["A4"].value == "Employee"


def test_selected_columns_local_times_durations_and_same_day_merge(
    scoped_client, db, monkeypatch
):
    from app.services import timesheet_report_service

    monkeypatch.setattr(
        timesheet_report_service, "utc_now", lambda: datetime(2026, 9, 22, 9)
    )
    add_session(
        db,
        start=datetime(2026, 9, 22, 6, 15),
        end=datetime(2026, 9, 22, 7, 15),
        work=99999,
        breaks=99999,
        late=15,
        transitions=[
            ("working", datetime(2026, 9, 22, 6, 15)),
            ("finished", datetime(2026, 9, 22, 7, 15)),
        ],
    )
    add_session(
        db,
        start=datetime(2026, 9, 22, 8),
        end=None,
        status=WorkStatus.BREAK,
        work=99999,
        breaks=99999,
        late=180,
        transitions=[
            ("working", datetime(2026, 9, 22, 8)),
            ("break", datetime(2026, 9, 22, 8, 30)),
        ],
    )
    db.commit()
    ws = sheet(
        export(
            scoped_client("admin"),
            columns=[
                "status",
                "end",
                "start",
                "work",
                "break",
                "lateness",
                "employee",
                "date",
            ],
        )
    )
    assert [cell.value for cell in ws[1]] == [
        "Статус",
        "Окончание",
        "Начало",
        "Работа",
        "Перерыв",
        "Опоздания",
        "Сотрудник",
        "Дата",
    ]
    assert ws["F2"].value == "Начало дня"
    assert ws["F3"].value == "Минуты"
    assert [cell.value for cell in ws[4]] == [
        "Перерыв",
        None,
        datetime(2026, 9, 22, 9, 15),
        timedelta(seconds=5400),
        timedelta(seconds=1800),
        195,
        "Employee",
        datetime(2026, 9, 22),
    ]
    assert [ws.cell(4, col).number_format for col in (3, 4, 5, 6, 8)] == [
        "dd.mm.yyyy hh:mm",
        "[h]:mm",
        "[h]:mm",
        "0",
        "dd.mm.yyyy",
    ]
    assert ws.freeze_panes == "A4"
    assert ws.auto_filter.ref == "A3:H4"


def test_zero_duration_is_numeric_excel_zero(scoped_client, db):
    add_session(db, work=0, breaks=0, late=0)
    db.commit()
    response = export(scoped_client("admin"), columns=["work", "break", "lateness"])
    ws = sheet(response)
    assert [cell.value for cell in ws[4]] == [timedelta(0), timedelta(0), 0]
    with ZipFile(BytesIO(response.content)) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for coordinate in ("A4", "B4", "C4"):
        cell = root.find(f".//x:c[@r='{coordinate}']", ns)
        assert cell is not None
        assert cell.get("t", "n") == "n"
        assert cell.find("x:v", ns).text == "0"


def test_distinct_group_timezones_and_status_labels(scoped_client, db):
    group = WidgetGroup(id=12, account_id=10, name="Berlin", timezone="Europe/Berlin")
    db.add(group)
    db.flush()
    db.add(
        User(
            id=5,
            amocrm_account_id=10,
            amocrm_user_id=105,
            name="Berlin employee",
            amocrm_rights={"is_admin": False},
        )
    )
    db.flush()
    db.add(
        GroupMember(
            id=5, account_id=10, group_id=12, user_id=5, is_active=True, track_time=True
        )
    )
    add_session(db, external_id=105, end=None, status=WorkStatus.WORKING)
    add_session(db, status=WorkStatus.FINISHED)
    db.commit()
    ws = sheet(
        export(scoped_client("admin"), columns=["employee", "start", "end", "status"])
    )
    values = {row[0]: row[1:] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert values["Employee"] == (
        datetime(2026, 9, 22, 9),
        datetime(2026, 9, 22, 10),
        "Закончил(а)",
    )
    assert values["Berlin employee"] == (datetime(2026, 9, 22, 8), None, "Работает")


def test_denial_and_foreign_filters_match_preview(scoped_client, db):
    add_session(db)
    db.commit()
    assert export(scoped_client("employee")).json()["error"]["code"] == "ACCESS_DENIED"
    for kwargs in ({"group_id": 20}, {"group_id": 11}, {"user_id": 4}):
        response = export(scoped_client("manager"), **kwargs)
        assert (response.status_code, response.json()["error"]["code"]) == (
            404,
            "NOT_FOUND",
        )


@pytest.mark.parametrize("actor", ["admin", "manager"])
def test_export_filter_uses_accessible_internal_user_id(scoped_client, db, actor):
    add_session(db)
    db.commit()
    ws = sheet(export(scoped_client(actor), user_id=2, columns=["employee", "work"]))
    assert ws.max_row == 2
    assert [cell.value for cell in ws[2]] == ["Employee", timedelta(hours=1)]


def test_employee_with_internal_user_filter_is_denied_before_visibility_lookup(
    scoped_client, db
):
    add_session(db)
    db.commit()
    response = export(scoped_client("employee"), user_id=2)
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "ACCESS_DENIED"


@pytest.mark.parametrize(
    "start,end,code",
    [
        ("2026-09-23", "2026-09-22", "REPORT_DATE_RANGE_INVALID"),
        ("2026-01-31", "2026-05-01", "REPORT_RANGE_LIMIT"),
        ("9999-12-31", "9999-12-31", "REPORT_DATE_OUT_OF_RANGE"),
    ],
)
def test_period_errors_keep_stable_russian_contract(scoped_client, start, end, code):
    response = export(scoped_client("admin"), date_from=start, date_to=end)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == code
    assert any("а" <= char <= "я" for char in response.json()["error"]["message"])


def test_forty_foreign_accounts_cannot_enter_export(scoped_client, db):
    add_session(db, work=77)
    for offset in range(40):
        account = 1000 + offset
        user_id = 1000 + offset
        db.add(
            User(
                id=user_id,
                amocrm_account_id=account,
                amocrm_user_id=102,
                name="PRIVATE foreign",
                amocrm_rights={"is_admin": False},
            )
        )
        db.add(
            WidgetGroup(
                id=user_id, account_id=account, name=f"Foreign {offset}", timezone="UTC"
            )
        )
        db.flush()
        db.add(
            GroupMember(
                id=user_id,
                account_id=account,
                group_id=user_id,
                user_id=user_id,
                is_active=True,
                track_time=True,
            )
        )
        add_session(db, account=account, work=999)
    db.commit()
    ws = sheet(export(scoped_client("admin")))
    assert ws.max_row == 4
    assert ws["A4"].value == "Employee"
    assert ws["F4"].value == timedelta(seconds=77)


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_formula_text_and_private_fields_never_become_workbook_content(
    scoped_client, db, prefix
):
    db.get(User, 2).name = f'  {prefix}HYPERLINK("https://evil.invalid")'
    add_session(db)
    db.commit()
    response = export(scoped_client("admin"))
    ws = sheet(response)
    assert ws["A4"].value == f'\'  {prefix}HYPERLINK("https://evil.invalid")'
    assert ws["A4"].data_type == "s"
    assert all(cell.hyperlink is None for row in ws for cell in row)
    with ZipFile(BytesIO(response.content)) as archive:
        xml = b"\n".join(
            archive.read(name) for name in archive.namelist() if name.endswith(".xml")
        )
    for secret in (b"PRIVATE", b"stale", b"amocrm", b"late_reason"):
        assert secret not in xml


def test_export_rejects_employee_day_count_above_limit(scoped_client, db, monkeypatch):
    from app.services import timesheet_report_service

    monkeypatch.setattr(timesheet_report_service, "MAX_EXPORT_ROWS", 1)
    add_session(db)
    add_session(
        db,
        day=date(2026, 9, 23),
        start=datetime(2026, 9, 23, 6),
        end=datetime(2026, 9, 23, 7),
    )
    db.commit()
    response = export(scoped_client("admin"), date_to="2026-09-23")
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "REPORT_EXPORT_TOO_LARGE"
    assert any("а" <= char <= "я" for char in response.json()["error"]["message"])


def test_real_ten_thousand_day_boundary(db):
    from app.services.timesheet_report_service import TimesheetReportService

    users = [
        User(
            id=100 + index,
            amocrm_account_id=10,
            amocrm_user_id=200 + index,
            name=f"Person {index:03d}",
            amocrm_rights={"is_admin": False},
        )
        for index in range(110)
    ]
    db.add_all(users)
    db.flush()
    db.add_all(
        GroupMember(
            id=100 + index,
            account_id=10,
            group_id=10,
            user_id=100 + index,
            is_active=True,
            track_time=True,
        )
        for index in range(110)
    )
    db.flush()

    def session_values(index):
        day = date(2026, 9, 1) + timedelta(days=index % 91)
        return dict(
            amocrm_account_id=10,
            amocrm_user_id=200 + index // 91,
            user_name="old snapshot",
            business_date=day,
            start_time=datetime.combine(day, datetime.min.time()),
            end_time=datetime.combine(day, datetime.min.time()) + timedelta(hours=1),
            current_status="finished",
            total_work_time=3600,
            total_break_time=0,
        )

    for start in range(0, 10000, 1000):
        db.execute(
            WorkSession.__table__.insert(),
            [session_values(index) for index in range(start, start + 1000)],
        )
    db.commit()
    service = TimesheetReportService(db)
    context = RequestContext(account_id=10, user=db.get(User, 1))
    assert (
        len(service.export_rows(context, date(2026, 9, 1), date(2026, 11, 30))) == 10000
    )
    db.execute(WorkSession.__table__.insert(), [session_values(10000)])
    db.commit()
    with pytest.raises(APIProblem) as error:
        service.export_rows(context, date(2026, 9, 1), date(2026, 11, 30))
    assert getattr(error.value, "status_code", None) == 413
    assert getattr(error.value, "code", None) == "REPORT_EXPORT_TOO_LARGE"


def test_export_query_count_is_bounded_without_n_plus_one(db):
    from app.services.timesheet_report_service import TimesheetReportService

    for offset in range(40):
        user_id = 100 + offset
        external_id = 200 + offset
        db.add(
            User(
                id=user_id,
                amocrm_account_id=10,
                amocrm_user_id=external_id,
                name=f"Person {offset}",
                amocrm_rights={"is_admin": False},
            )
        )
        db.flush()
        db.add(
            GroupMember(
                id=user_id,
                account_id=10,
                group_id=10,
                user_id=user_id,
                is_active=True,
                track_time=True,
            )
        )
        add_session(db, external_id=external_id)
    db.commit()
    statements = []

    def record(_conn, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)

    event.listen(db.get_bind(), "before_cursor_execute", record)
    try:
        rows = TimesheetReportService(db).export_rows(
            RequestContext(account_id=10, user=db.get(User, 1)),
            date(2026, 9, 22),
            date(2026, 9, 22),
        )
    finally:
        event.remove(db.get_bind(), "before_cursor_execute", record)
    assert len(rows) == 40
    assert len(statements) <= 6, repr(statements)


def test_export_fetches_employee_days_in_bounded_batches(db, monkeypatch):
    from app.services import timesheet_report_service
    from app.services.timesheet_report_service import TimesheetReportService

    monkeypatch.setattr(timesheet_report_service, "REPORT_ROW_BATCH_SIZE", 3)
    for offset in range(8):
        user_id = 100 + offset
        external_id = 200 + offset
        db.add(
            User(
                id=user_id,
                amocrm_account_id=10,
                amocrm_user_id=external_id,
                name=f"Person {offset}",
                amocrm_rights={"is_admin": False},
            )
        )
        db.flush()
        db.add(
            GroupMember(
                id=user_id,
                account_id=10,
                group_id=10,
                user_id=user_id,
                is_active=True,
                track_time=True,
            )
        )
        add_session(db, external_id=external_id)
    db.commit()

    service = TimesheetReportService(db)
    original = service._rows_for_keys
    batch_sizes = []

    def record_batch(*args, **kwargs):
        ordered_keys = args[4]
        batch_sizes.append(len(ordered_keys))
        return original(*args, **kwargs)

    monkeypatch.setattr(service, "_rows_for_keys", record_batch)
    rows = service.export_rows(
        RequestContext(account_id=10, user=db.get(User, 1)),
        date(2026, 9, 22),
        date(2026, 9, 22),
    )

    assert batch_sizes == [3, 3, 2]
    assert [row.employee_name for row in rows] == [
        f"Person {offset}" for offset in range(8)
    ]


def test_employee_day_details_reject_unbounded_session_source(db, monkeypatch):
    from app.services import timesheet_report_service
    from app.services.timesheet_report_service import TimesheetReportService

    add_session(db, start=datetime(2026, 9, 22, 6), end=datetime(2026, 9, 22, 7))
    add_session(db, start=datetime(2026, 9, 22, 8), end=datetime(2026, 9, 22, 9))
    db.commit()
    visible = {2: (db.get(User, 2), db.get(WidgetGroup, 10))}
    keys = [SimpleNamespace(id=2, business_date=date(2026, 9, 22))]

    monkeypatch.setattr(timesheet_report_service, "REPORT_SOURCE_ROW_LIMIT", 1)
    with pytest.raises(APIProblem) as error:
        TimesheetReportService(db)._rows_for_keys(
            10,
            visible,
            date(2026, 9, 22),
            date(2026, 9, 22),
            keys,
            datetime(2026, 9, 22, 10),
        )
    assert error.value.status_code == 413
    assert error.value.code == "REPORT_EXPORT_SOURCE_TOO_LARGE"


def test_employee_day_details_reject_unbounded_transition_source(db, monkeypatch):
    from app.services import timesheet_report_service
    from app.services.timesheet_report_service import TimesheetReportService

    add_session(
        db,
        start=datetime(2026, 9, 22, 6),
        end=datetime(2026, 9, 22, 7),
        transitions=[
            ("break", datetime(2026, 9, 22, 6, 15)),
            ("working", datetime(2026, 9, 22, 6, 30)),
        ],
    )
    db.commit()
    visible = {2: (db.get(User, 2), db.get(WidgetGroup, 10))}
    keys = [SimpleNamespace(id=2, business_date=date(2026, 9, 22))]

    monkeypatch.setattr(timesheet_report_service, "REPORT_SOURCE_ROW_LIMIT", 1)
    with pytest.raises(APIProblem) as error:
        TimesheetReportService(db)._rows_for_keys(
            10,
            visible,
            date(2026, 9, 22),
            date(2026, 9, 22),
            keys,
            datetime(2026, 9, 22, 8),
        )
    assert error.value.status_code == 413
    assert error.value.code == "REPORT_EXPORT_SOURCE_TOO_LARGE"


@pytest.fixture
def postgres_db():
    admin_url = os.getenv("TEST_POSTGRES_ADMIN_URL")
    if not admin_url:
        pytest.skip("TEST_POSTGRES_ADMIN_URL required for disposable PostgreSQL test")
    url = make_url(admin_url)
    assert url.get_backend_name() == "postgresql"
    name = "timesheet_excel_test_" + uuid4().hex
    admin = create_engine(url, isolation_level="AUTOCOMMIT")
    with admin.connect() as connection:
        connection.execute(text(f'CREATE DATABASE "{name}"'))
    engine = create_engine(url.set(database=name))
    Base.metadata.create_all(engine)
    database = sessionmaker(bind=engine)()
    try:
        yield database
    finally:
        database.close()
        engine.dispose()
        with admin.connect() as connection:
            connection.execute(text(f'DROP DATABASE "{name}" WITH (FORCE)'))
        admin.dispose()


def test_postgres_export_uses_scoped_canonical_rows(postgres_db):
    from app.services.timesheet_report_service import TimesheetReportService
    from app.services.timesheet_excel_service import TimesheetExcelService
    from app.schemas.report import TimesheetColumn

    db = postgres_db
    admin = User(
        id=1,
        amocrm_user_id=101,
        amocrm_account_id=10,
        name="Admin",
        amocrm_rights={"is_admin": True},
    )
    employee = User(
        id=2,
        amocrm_user_id=102,
        amocrm_account_id=10,
        name="Employee",
        amocrm_rights={"is_admin": False},
    )
    db.add_all(
        [
            admin,
            employee,
            WidgetGroup(id=10, account_id=10, name="Team", timezone="Europe/Minsk"),
        ]
    )
    db.flush()
    db.add(
        GroupMember(
            id=1, account_id=10, group_id=10, user_id=2, is_active=True, track_time=True
        )
    )
    add_session(db, work=77)
    db.commit()
    rows = TimesheetReportService(db).export_rows(
        RequestContext(account_id=10, user=admin), date(2026, 9, 22), date(2026, 9, 22)
    )
    assert len(rows) == 1
    assert (rows[0].employee_name, rows[0].work_seconds, rows[0].group_timezone) == (
        "Employee",
        77,
        "Europe/Minsk",
    )
    content = TimesheetExcelService().render(
        rows, [TimesheetColumn.EMPLOYEE, TimesheetColumn.WORK]
    )
    ws = load_workbook(BytesIO(content)).active
    assert [cell.value for cell in ws[2]] == ["Employee", timedelta(seconds=77)]
