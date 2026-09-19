from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base
from app.models.user import User
from app.models.user import UserRole
from app.models.work_session import WorkSession
from app.models.activity_session import ActivitySession, EntityType
from app.models.report import Report, ReportType, ReportFormat
from app.models.department import Department
from app.models.widget_group import WidgetGroup
from app.models.group_member import GroupMember
from datetime import datetime
import pytest


def _client(monkeypatch):
    from app.core.database import get_db
    from app.main import app

    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add_all(
        [
            User(id=1, amocrm_user_id=10, amocrm_account_id=20, name="Self"),
            User(id=2, amocrm_user_id=11, amocrm_account_id=20, name="Other"),
        ]
    )
    db.commit()
    monkeypatch.setenv("ENVIRONMENT", "test")
    # Each test owns fresh middleware state, including the real rate limiter.
    monkeypatch.setattr(app, "middleware_stack", None)
    app.dependency_overrides[get_db] = lambda: db
    return TestClient(app), app


def test_settings_rejects_a_forged_account_path_without_existence_leak(monkeypatch):
    """Using path account_id instead of RequestContext must fail this authorization test."""
    client, app = _client(monkeypatch)
    try:
        response = client.get(
            "/api/v1/settings/21", headers={"X-User-Id": "10", "X-Account-Id": "20"}
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_session_mutation_rejects_another_employee_before_service_execution(
    monkeypatch,
):
    """Replacing self-only target validation with a service call would expose mutation."""
    client, app = _client(monkeypatch)
    try:
        response = client.post(
            "/api/v1/sessions/break/11",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_rate_limit_uses_the_standard_429_error_body():
    """Returning a legacy detail-only rate-limit response must fail this API contract."""
    from app.main import RateLimitMiddleware

    app = FastAPI()
    app.add_middleware(RateLimitMiddleware, calls_per_minute=1)

    @app.get("/")
    def root():
        return {"ok": True}

    client = TestClient(app)
    assert client.get("/").status_code == 200
    response = client.get("/")
    assert response.status_code == 429
    assert response.json()["error"]["code"] == "RATE_LIMITED"


def test_duplicate_session_returns_normalized_conflict(monkeypatch):
    """Mapping a session state conflict to 400 would hide a retryable conflict from clients."""
    client, app = _client(monkeypatch)
    headers = {"X-User-Id": "10", "X-Account-Id": "20"}
    try:
        assert (
            client.post(
                "/api/v1/sessions/start",
                json={"user_id": 10, "user_name": "Self"},
                headers=headers,
            ).status_code
            == 201
        )
        response = client.post(
            "/api/v1/sessions/start",
            json={"user_id": 10, "user_name": "Self"},
            headers=headers,
        )
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "SESSION_CONFLICT"
    finally:
        app.dependency_overrides.clear()


def test_activity_work_session_path_rejects_foreign_session_before_handler(monkeypatch):
    """Activity routes must scope work_session_id, not only the /sessions routes."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(
        WorkSession(
            id=40,
            amocrm_user_id=11,
            amocrm_account_id=20,
            user_name="Other",
            start_time=datetime.utcnow(),
        )
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/activity/current/40",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_activity_session_path_rejects_foreign_session_before_handler(monkeypatch):
    """Activity session IDs inherit ownership from their work session."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(
        WorkSession(
            id=41,
            amocrm_user_id=11,
            amocrm_account_id=20,
            user_name="Other",
            start_time=datetime.utcnow(),
        )
    )
    db.add(
        ActivitySession(
            id=42,
            work_session_id=41,
            entity_type=EntityType.LEAD,
            entity_id=7,
            start_time=datetime.utcnow(),
        )
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/activity/events/42",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_report_path_rejects_foreign_report_before_handler(monkeypatch):
    """Saved reports must be account-scoped by their own account column."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(
        Report(
            id=50,
            report_type=ReportType.DAILY,
            report_format=ReportFormat.JSON,
            title="Other account",
            start_date=datetime.utcnow(),
            end_date=datetime.utcnow(),
            account_id="21",
            data={},
            generated_by=11,
        )
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/reports/50",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_department_path_rejects_unowned_department_before_handler(monkeypatch):
    """Department IDs without an active user in this account are not addressable."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(
        Department(
            id=60,
            name="Other department",
            work_start_time=datetime.strptime("09:00", "%H:%M").time(),
            work_end_time=datetime.strptime("18:00", "%H:%M").time(),
        )
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/departments/60/schedule",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_excel_body_rejects_unowned_department_before_handler(monkeypatch):
    """JSON body department IDs must receive the same scope check as path IDs."""
    client, app = _client(monkeypatch)
    try:
        response = client.post(
            "/api/v1/excel/department",
            json={
                "date_from": "2026-09-01",
                "date_to": "2026-09-18",
                "department_ids": [999],
            },
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_activity_history_limit_is_bounded(monkeypatch):
    """Legacy activity history must not accept an unbounded item request."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(
        WorkSession(
            id=70,
            amocrm_user_id=10,
            amocrm_account_id=20,
            user_name="Self",
            start_time=datetime.utcnow(),
        )
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/activity/history/70?limit=1001",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 422
    finally:
        app.dependency_overrides.clear()


def test_id_free_team_collections_filter_to_employee_self_and_account(monkeypatch):
    """A collection without IDs must still apply the verified user scope."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    now = datetime.utcnow()
    db.add_all(
        [
            WorkSession(
                id=80,
                amocrm_user_id=10,
                amocrm_account_id=20,
                user_name="Self",
                start_time=now,
                total_work_time=5,
            ),
            WorkSession(
                id=81,
                amocrm_user_id=11,
                amocrm_account_id=20,
                user_name="Other",
                start_time=now,
                total_work_time=50,
            ),
            WorkSession(
                id=82,
                amocrm_user_id=99,
                amocrm_account_id=21,
                user_name="Foreign account",
                start_time=now,
                total_work_time=500,
            ),
        ]
    )
    db.commit()
    headers = {"X-User-Id": "10", "X-Account-Id": "20"}
    try:
        stats = client.get("/api/v1/team/stats", headers=headers)
        assert stats.status_code == 200
        assert stats.json()["total_members"] == 1
        assert stats.json()["total_work_time"] == 5

        activity = client.get("/api/v1/team/activity", headers=headers)
        assert activity.status_code == 200
        assert [row["user_id"] for row in activity.json()] == [10]
    finally:
        app.dependency_overrides.clear()


def test_id_free_team_collections_filter_manager_to_own_group_and_self(monkeypatch):
    """A manager sees self and active members of own group, never another group."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    manager = User(
        id=3,
        amocrm_user_id=12,
        amocrm_account_id=20,
        name="Manager",
        role=UserRole.ROP,
        amocrm_role_id=77,
        amocrm_rights={"role_id": 77, "is_admin": False},
    )
    foreign = User(
        id=4,
        amocrm_user_id=13,
        amocrm_account_id=20,
        name="Foreign group",
    )
    own_group = WidgetGroup(
        id=30, account_id=20, name="Own", manager_user_id=3, manager_role_id=77
    )
    other_group = WidgetGroup(
        id=31, account_id=20, name="Other", manager_user_id=4, manager_role_id=88
    )
    db.add_all([manager, foreign, own_group, other_group])
    db.flush()
    db.add_all(
        [
            GroupMember(account_id=20, group_id=30, user_id=2, is_active=True),
            GroupMember(account_id=20, group_id=31, user_id=4, is_active=True),
            WorkSession(
                id=83,
                amocrm_user_id=12,
                amocrm_account_id=20,
                user_name="Manager",
                start_time=datetime.utcnow(),
                total_work_time=7,
            ),
            WorkSession(
                id=84,
                amocrm_user_id=11,
                amocrm_account_id=20,
                user_name="Other",
                start_time=datetime.utcnow(),
                total_work_time=8,
            ),
            WorkSession(
                id=85,
                amocrm_user_id=13,
                amocrm_account_id=20,
                user_name="Foreign group",
                start_time=datetime.utcnow(),
                total_work_time=80,
            ),
        ]
    )
    db.commit()
    try:
        response = client.get(
            "/api/v1/team/stats",
            headers={"X-User-Id": "12", "X-Account-Id": "20"},
        )
        assert response.status_code == 200
        assert response.json()["total_members"] == 2
        assert response.json()["total_work_time"] == 15
    finally:
        app.dependency_overrides.clear()


def test_account_report_without_user_filter_is_visibility_scoped(monkeypatch):
    """Account-wide report queries must not turn omitted user_id into all users."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    day = datetime.utcnow().replace(hour=12, minute=0, second=0, microsecond=0)
    db.add_all(
        [
            WorkSession(
                id=90,
                amocrm_user_id=10,
                amocrm_account_id=20,
                user_name="Self",
                start_time=day,
                total_work_time=5,
            ),
            WorkSession(
                id=91,
                amocrm_user_id=11,
                amocrm_account_id=20,
                user_name="Other",
                start_time=day,
                total_work_time=50,
            ),
        ]
    )
    db.commit()
    try:
        response = client.get(
            f"/api/v1/reports/daily?account_id=20&date={day.date().isoformat()}",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["total_users"] == 1
        assert [row["user_id"] for row in payload["sessions"]] == [10]
    finally:
        app.dependency_overrides.clear()


def test_id_free_excel_collections_filter_to_verified_account(monkeypatch):
    """Excel aggregate exports must not include rows from another account."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    day = datetime.utcnow()
    department = Department(
        id=92,
        name="Shared legacy department",
        work_start_time=datetime.strptime("09:00", "%H:%M").time(),
        work_end_time=datetime.strptime("18:00", "%H:%M").time(),
    )
    db.get(User, 1).department_id = 92
    db.add(
        User(
            id=99,
            amocrm_user_id=99,
            amocrm_account_id=21,
            name="Foreign account",
            department_id=92,
        )
    )
    db.add_all(
        [
            WorkSession(
                id=92,
                amocrm_user_id=10,
                amocrm_account_id=20,
                user_name="Self",
                start_time=day,
                total_work_time=5,
                is_late=True,
                late_minutes=3,
            ),
            WorkSession(
                id=93,
                amocrm_user_id=99,
                amocrm_account_id=21,
                user_name="Foreign account",
                start_time=day,
                total_work_time=500,
                is_late=True,
                late_minutes=30,
            ),
        ]
    )
    db.add(department)
    db.commit()
    headers = {"X-User-Id": "10", "X-Account-Id": "20"}
    payload = {
        "date_from": day.date().isoformat(),
        "date_to": day.date().isoformat(),
    }
    try:
        # Admin is required by the endpoint when no department filter is given.
        db.get(User, 1).role = UserRole.ADMIN
        db.get(User, 1).amocrm_rights = {"is_admin": True}
        db.commit()
        department = client.post(
            "/api/v1/excel/department", json=payload, headers=headers
        )
        late = client.post("/api/v1/excel/late-arrivals", json=payload, headers=headers)
        assert department.status_code == 200
        assert late.status_code == 200

        from openpyxl import load_workbook
        from io import BytesIO

        department_sheet = load_workbook(
            BytesIO(department.content), read_only=True
        ).active
        late_sheet = load_workbook(BytesIO(late.content), read_only=True).active
        assert all(
            "Foreign account" not in str(cell.value)
            for row in department_sheet.iter_rows()
            for cell in row
        )
        assert all(
            "Foreign account" not in str(cell.value)
            for row in late_sheet.iter_rows()
            for cell in row
        )
    finally:
        app.dependency_overrides.clear()


def test_user_id_semantics_are_route_specific_when_internal_and_external_collide(
    monkeypatch,
):
    """Sessions use external IDs; KPI/Excel targets use internal IDs."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    # Internal id 10 belongs to a different user than external amoCRM id 10.
    db.add(User(id=10, amocrm_user_id=100, amocrm_account_id=20, name="Collision"))
    db.commit()
    headers = {"X-User-Id": "10", "X-Account-Id": "20"}
    try:
        session = client.get("/api/v1/sessions/current/10", headers=headers)
        assert session.status_code == 200
        assert session.json() is None

        kpi = client.get("/api/v1/kpi/user/10", headers=headers)
        assert kpi.status_code == 404
        assert kpi.json()["error"]["code"] == "NOT_FOUND"

        excel = client.post(
            "/api/v1/excel/employee/10",
            json={"date_from": "2026-09-01", "date_to": "2026-09-18"},
            headers=headers,
        )
        assert excel.status_code == 404
        assert excel.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_missing_context_returns_normalized_401(monkeypatch):
    client, app = _client(monkeypatch)
    try:
        response = client.get("/api/v1/settings/20")
        assert response.status_code == 401
        assert response.json()["error"]["code"] == "AMOCRM_TOKEN_EXPIRED"
    finally:
        app.dependency_overrides.clear()


def test_employee_forbidden_operation_returns_normalized_403(monkeypatch):
    client, app = _client(monkeypatch)
    try:
        response = client.post(
            "/api/v1/excel/department",
            json={"date_from": "2026-09-01", "date_to": "2026-09-18"},
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 403
        assert response.json()["error"]["code"] == "ACCESS_DENIED"
    finally:
        app.dependency_overrides.clear()


def test_category_account_owner_allows_first_event_and_denies_foreign_category(
    monkeypatch,
):
    """Ownership comes from migration-007 account_id, not an existing event."""
    client, app = _client(monkeypatch)
    from app.core.database import get_db
    from app.models.activity_category import ActivityCategory

    db = app.dependency_overrides[get_db]()
    db.add(
        WorkSession(
            id=100,
            amocrm_user_id=10,
            amocrm_account_id=20,
            user_name="Self",
            start_time=datetime.utcnow(),
        )
    )
    db.add(
        ActivitySession(
            id=101,
            work_session_id=100,
            entity_type=EntityType.LEAD,
            entity_id=1,
            start_time=datetime.utcnow(),
        )
    )
    own = ActivityCategory(
        id=102,
        account_id=20,
        name="own",
        display_name="Own",
        color="#fff",
    )
    foreign = ActivityCategory(
        id=103,
        account_id=21,
        name="foreign",
        display_name="Foreign",
        color="#000",
    )
    db.add_all([own, foreign])
    db.commit()
    headers = {"X-User-Id": "10", "X-Account-Id": "20"}
    try:
        first_event = client.post(
            "/api/v1/activity/event?activity_session_id=101&event_type=card_opened&category_id=102",
            headers=headers,
        )
        assert first_event.status_code == 201

        foreign_response = client.get("/api/v1/categories/103", headers=headers)
        assert foreign_response.status_code == 404
        assert foreign_response.json()["error"]["code"] == "NOT_FOUND"
    finally:
        app.dependency_overrides.clear()


def test_categories_are_account_scoped_and_set_display_name(monkeypatch):
    client, app = _client(monkeypatch)
    from app.core.database import get_db

    db = app.dependency_overrides[get_db]()
    db.add(User(id=3, amocrm_user_id=30, amocrm_account_id=21, name="Other account"))
    db.commit()
    payload = {"name": "shared", "display_name": "Shared", "color": "#fff"}
    try:
        own = client.post(
            "/api/v1/categories?account_id=20",
            json=payload,
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        foreign = client.post(
            "/api/v1/categories?account_id=21",
            json=payload,
            headers={"X-User-Id": "30", "X-Account-Id": "21"},
        )
        assert own.status_code == 201
        assert own.json()["display_name"] == "Shared"
        assert foreign.status_code == 201
        assert foreign.json()["account_id"] == 21
    finally:
        app.dependency_overrides.clear()


@pytest.fixture
def scoped_client(monkeypatch):
    """Two assigned groups have different snapshots; only role 78 is live."""
    from app.core.database import get_db

    client, app = _client(monkeypatch)
    db = app.dependency_overrides[get_db]()
    db.add_all(
        [
            User(
                id=3,
                amocrm_user_id=12,
                amocrm_account_id=20,
                name="Manager",
                amocrm_role_id=78,
                amocrm_rights={"role_id": 78, "is_admin": False},
            ),
            User(
                id=4,
                amocrm_user_id=13,
                amocrm_account_id=20,
                name="Admin",
                amocrm_rights={"is_admin": True},
            ),
            User(id=5, amocrm_user_id=14, amocrm_account_id=21, name="Foreign"),
            WidgetGroup(
                id=30,
                account_id=20,
                name="Current",
                manager_user_id=3,
                manager_role_id=78,
            ),
            WidgetGroup(
                id=31,
                account_id=20,
                name="Stale",
                manager_user_id=3,
                manager_role_id=77,
            ),
            GroupMember(account_id=20, group_id=30, user_id=1),
            GroupMember(account_id=20, group_id=31, user_id=2),
        ]
    )
    for dept_id, user_id in [(60, 1), (61, 2), (62, 5)]:
        db.add(
            Department(
                id=dept_id,
                name=f"Department {dept_id}",
                work_start_time=datetime.strptime("09:00", "%H:%M").time(),
                work_end_time=datetime.strptime("18:00", "%H:%M").time(),
            )
        )
        db.get(User, user_id).department_id = dept_id
    for session_id, external_id, account in [
        (100, 10, 20),
        (101, 11, 20),
        (102, 14, 21),
    ]:
        db.add(
            WorkSession(
                id=session_id,
                amocrm_user_id=external_id,
                amocrm_account_id=account,
                user_name=f"User {external_id}",
                start_time=datetime(2026, 9, 18, 10),
                total_work_time=10,
            )
        )
    db.commit()
    try:
        yield client, db
    finally:
        app.dependency_overrides.clear()
        db.close()


@pytest.mark.parametrize(
    "actor,department,expected",
    [
        (10, 60, 200),
        (10, 61, 404),
        (10, 62, 404),
        (12, 60, 200),
        (12, 61, 404),
        (12, 62, 404),
        (13, 60, 200),
        (13, 61, 200),
        (13, 62, 404),
    ],
)
def test_department_schedule_requires_user_visibility(
    scoped_client, actor, department, expected
):
    client, _ = scoped_client
    response = client.get(
        f"/api/v1/departments/{department}/schedule",
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == expected
    if expected == 404:
        assert response.json()["error"]["code"] == "NOT_FOUND"


@pytest.mark.parametrize("actor,expected", [(10, 403), (12, [60]), (13, [60, 61])])
def test_department_collection_scope(scoped_client, actor, expected):
    client, _ = scoped_client
    response = client.get(
        "/api/v1/departments/",
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    if isinstance(expected, int):
        assert response.status_code == expected
    else:
        assert response.status_code == 200
        assert sorted(row["id"] for row in response.json()) == expected


@pytest.mark.parametrize(
    "actor,department,expected",
    [(10, 60, 403), (12, 60, 403), (13, 60, 200), (13, 62, 404)],
)
def test_department_schedule_mutation_requires_account_admin(
    scoped_client, actor, department, expected
):
    client, _ = scoped_client
    response = client.put(
        f"/api/v1/departments/{department}/schedule",
        json={"work_start_time": "08:00:00"},
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == expected


@pytest.mark.parametrize("actor", [12, 13])
@pytest.mark.parametrize("source", ["query", "body"])
def test_report_cannot_be_attributed_to_visible_subordinate(
    scoped_client, actor, source
):
    client, db = scoped_client
    payload = {
        "report_type": "custom",
        "start_date": "2026-09-18",
        "end_date": "2026-09-18",
    }
    author = 3 if actor == 12 else 4
    claimed_author = 1 if source == "query" else author
    if source == "body":
        payload["generated_by"] = 1
    response = client.post(
        f"/api/v1/reports/generate?account_id=20&generated_by={claimed_author}",
        json=payload,
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "NOT_FOUND"
    assert db.query(Report).count() == 0


@pytest.mark.parametrize("actor,author", [(10, 1), (12, 3), (13, 4)])
def test_report_persists_verified_author(scoped_client, actor, author):
    client, db = scoped_client
    response = client.post(
        f"/api/v1/reports/generate?account_id=20&generated_by={author}",
        json={
            "report_type": "custom",
            "start_date": "2026-09-18",
            "end_date": "2026-09-18",
        },
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == 201
    assert response.json()["generated_by"] == author
    assert db.query(Report).one().generated_by == author


@pytest.mark.parametrize(
    "actor,target,session_id,expected",
    [
        (10, 10, 100, 200),
        (10, 11, 101, 404),
        (12, 10, 100, 200),
        (12, 11, 101, 404),
        (12, 14, 102, 404),
        (13, 11, 101, 200),
        (13, 14, 102, 404),
    ],
)
def test_session_reads_use_verified_visibility(
    scoped_client, actor, target, session_id, expected
):
    client, _ = scoped_client
    for path in [f"current/{target}", f"history/{target}", str(session_id)]:
        response = client.get(
            f"/api/v1/sessions/{path}",
            headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
        )
        assert response.status_code == expected, path
        if expected == 200:
            result = response.json()
            assert (result[0] if isinstance(result, list) else result)[
                "id"
            ] == session_id
        else:
            assert response.json()["error"]["code"] == "NOT_FOUND"


def test_session_reads_resolve_duplicate_external_id_in_verified_account(scoped_client):
    client, db = scoped_client
    db.add(
        User(
            id=6, amocrm_user_id=10, amocrm_account_id=21, name="Duplicate external ID"
        )
    )
    db.add(
        WorkSession(
            id=103,
            amocrm_user_id=10,
            amocrm_account_id=21,
            user_name="Duplicate external ID",
            start_time=datetime(2026, 9, 18, 11),
        )
    )
    db.commit()
    for path in ["current/10", "history/10", "100"]:
        response = client.get(
            f"/api/v1/sessions/{path}",
            headers={"X-User-Id": "10", "X-Account-Id": "20"},
        )
        assert response.status_code == 200
        result = response.json()
        assert (result[0] if isinstance(result, list) else result)["id"] == 100


@pytest.mark.parametrize(
    "actor,author,target,expected",
    [
        (10, 1, 10, 201),
        (10, 1, 11, 404),
        (12, 3, 10, 201),
        (12, 3, 11, 404),
        (13, 4, 11, 201),
        (13, 4, 14, 404),
    ],
)
def test_employee_report_generation_preserves_visibility(
    scoped_client, actor, author, target, expected
):
    client, db = scoped_client
    response = client.post(
        f"/api/v1/reports/generate?account_id=20&generated_by={author}",
        json={
            "report_type": "employee",
            "user_id": target,
            "start_date": "2026-09-18",
            "end_date": "2026-09-18",
        },
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == expected
    if expected == 201:
        saved = db.query(Report).one()
        assert saved.generated_by == author
        assert saved.data["user_id"] == target
        assert [row["user_id"] for row in saved.data["sessions"]] == [target]
    else:
        assert db.query(Report).count() == 0


def test_employee_report_generation_without_sessions_returns_not_found(scoped_client):
    client, db = scoped_client
    response = client.post(
        "/api/v1/reports/generate?account_id=20&generated_by=1",
        json={
            "report_type": "employee",
            "user_id": 10,
            "start_date": "2026-09-01",
            "end_date": "2026-09-01",
        },
        headers={"X-User-Id": "10", "X-Account-Id": "20"},
    )
    assert response.status_code == 404
    assert db.query(Report).count() == 0


def test_employee_report_service_rejects_user_outside_explicit_visibility(
    scoped_client,
):
    from app.services.report_service import ReportService
    from datetime import date

    _, db = scoped_client
    assert (
        ReportService.get_employee_report(
            db,
            "20",
            11,
            date(2026, 9, 18),
            date(2026, 9, 18),
            visible_external_user_ids={10},
        )
        is None
    )


def test_department_kpi_filters_members_by_group_not_legacy_department(scoped_client):
    client, db = scoped_client
    db.get(User, 2).department_id = 60
    db.commit()
    response = client.get(
        "/api/v1/kpi/department/60",
        headers={"X-User-Id": "12", "X-Account-Id": "20"},
    )
    assert response.status_code == 200
    assert response.json()["total_employees"] == 1


def test_manager_collections_exclude_members_of_stale_assignment(scoped_client):
    client, _ = scoped_client
    headers = {"X-User-Id": "12", "X-Account-Id": "20"}
    response = client.get("/api/v1/team/status", headers=headers)
    assert response.status_code == 200
    assert sorted(row["user_id"] for row in response.json()) == [10, 12]
    report = client.get(
        "/api/v1/reports/daily?account_id=20&date=2026-09-18", headers=headers
    )
    assert report.status_code == 200
    assert [row["user_id"] for row in report.json()["sessions"]] == [10]


@pytest.mark.parametrize("actor,expected", [(10, 403), (12, 403), (13, 201)])
def test_department_creation_requires_verified_admin(scoped_client, actor, expected):
    client, db = scoped_client
    response = client.post(
        "/api/v1/departments/",
        json={
            "name": "New department",
            "work_start_time": "09:00:00",
            "work_end_time": "18:00:00",
        },
        headers={"X-User-Id": str(actor), "X-Account-Id": "20"},
    )
    assert response.status_code == expected
    assert (db.query(Department).filter_by(name="New department").count() == 1) is (
        expected == 201
    )


@pytest.mark.parametrize("path", ["department", "late-arrivals"])
def test_department_excel_filters_group_members_within_shared_department(
    scoped_client, path
):
    from io import BytesIO
    from openpyxl import load_workbook

    client, db = scoped_client
    db.get(User, 2).department_id = 60
    for session_id in (100, 101):
        db.get(WorkSession, session_id).is_late = True
        db.get(WorkSession, session_id).late_minutes = 3
    db.commit()
    response = client.post(
        f"/api/v1/excel/{path}",
        json={
            "date_from": "2026-09-18",
            "date_to": "2026-09-18",
            "department_ids": [60],
        },
        headers={"X-User-Id": "12", "X-Account-Id": "20"},
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    cells = [cell.value for row in workbook.active.iter_rows() for cell in row]
    assert "Self" in cells
    assert "Other" not in cells
