"""Phase-9 backend acceptance stories against a migrated PostgreSQL database."""

from __future__ import annotations

import logging
import os
from datetime import datetime, time
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import GroupMember, User, WidgetGroup, WorkSession
from app.models.user import UserRole


ACCOUNT = 910
DAY = "2026-09-30"


def _user(
    id_: int,
    external_id: int,
    account_id: int,
    name: str,
    *,
    admin: bool = False,
    role_id: int = 1,
) -> User:
    return User(
        id=id_,
        amocrm_user_id=external_id,
        amocrm_account_id=account_id,
        name=name,
        role=UserRole.ADMIN if admin else UserRole.EMPLOYEE,
        amocrm_rights={"is_admin": admin, "role_id": role_id},
        amocrm_role_id=role_id,
        is_active=True,
    )


@pytest.fixture
def postgres_engine(monkeypatch):
    from app.core.config import settings

    admin_url = os.getenv("TEST_POSTGRES_ADMIN_URL")
    if not admin_url:
        pytest.skip("TEST_POSTGRES_ADMIN_URL required for backend acceptance")
    url = make_url(admin_url)
    assert url.get_backend_name() == "postgresql"
    name = "timesheet_phase9_e2e_" + uuid4().hex
    admin = create_engine(url, isolation_level="AUTOCOMMIT")
    with admin.connect() as connection:
        connection.execute(text(f'CREATE DATABASE "{name}"'))
    database_url = url.set(database=name)
    engine = create_engine(database_url)
    monkeypatch.setattr(
        settings,
        "DATABASE_URL",
        database_url.render_as_string(hide_password=False),
    )
    config = Config(str(Path(__file__).resolve().parents[2] / "alembic.ini"))
    config.set_main_option(
        "script_location", str(Path(__file__).resolve().parents[2] / "migrations")
    )
    try:
        command.upgrade(config, "head")
        yield engine
    finally:
        engine.dispose()
        with admin.connect() as connection:
            connection.execute(text(f'DROP DATABASE "{name}" WITH (FORCE)'))
        admin.dispose()


@pytest.fixture
def acceptance(postgres_engine, monkeypatch):
    """One migrated tenant plus 40 same-ID tenants proves all story boundaries."""
    from app.core.rate_limit import (
        authenticated_request_limiter,
        authentication_concurrency_guard,
        unknown_request_limiter,
    )
    from app.main import app

    db = Session(postgres_engine)
    admin = _user(1, 101, ACCOUNT, "Admin", admin=True, role_id=70)
    manager = _user(2, 102, ACCOUNT, "Manager", role_id=77)
    employee = _user(3, 103, ACCOUNT, "Employee", role_id=71)
    outsider = _user(4, 104, ACCOUNT, "Other group", role_id=72)
    own_group = WidgetGroup(
        id=10,
        account_id=ACCOUNT,
        name="Sales",
        timezone="UTC",
        work_start_time=time(9),
        work_end_time=time(18),
        manager_user_id=manager.id,
        manager_role_id=manager.amocrm_role_id,
        allow_restart_session=False,
    )
    other_group = WidgetGroup(
        id=20,
        account_id=ACCOUNT,
        name="Private",
        timezone="UTC",
    )
    db.add_all([admin, manager, employee, outsider])
    db.flush()
    db.add_all([own_group, other_group])
    db.flush()
    db.add_all(
        [
            GroupMember(
                account_id=ACCOUNT,
                group_id=10,
                user_id=3,
                is_active=True,
                track_time=True,
                hide_widget=True,
            ),
            GroupMember(
                account_id=ACCOUNT,
                group_id=20,
                user_id=4,
                is_active=True,
                track_time=True,
            ),
        ]
    )
    # Repeated external employee ID is intentional: every public story must
    # stay bound to the verified tenant even at the required 40-account scale.
    for offset in range(40):
        account_id = 10_000 + offset
        internal_id = 10_000 + offset
        db.add(_user(internal_id, 103, account_id, f"Foreign {offset}"))
        db.add(
            WidgetGroup(
                id=internal_id,
                account_id=account_id,
                name=f"Foreign {offset}",
                timezone="UTC",
            )
        )
        db.flush()
        db.add(
            GroupMember(
                account_id=account_id,
                group_id=internal_id,
                user_id=internal_id,
                is_active=True,
                track_time=True,
            )
        )
    db.commit()

    monkeypatch.setenv("ENVIRONMENT", "test")
    monkeypatch.setattr(app, "middleware_stack", None)
    app.dependency_overrides[get_db] = lambda: db
    authenticated_request_limiter.reset()
    unknown_request_limiter.reset()
    authentication_concurrency_guard.reset()
    with TestClient(app, raise_server_exceptions=False) as client:
        yield client, db
    app.dependency_overrides.clear()
    db.close()
    authenticated_request_limiter.reset()
    unknown_request_limiter.reset()
    authentication_concurrency_guard.reset()


def _headers(external_user_id: int, account_id: int = ACCOUNT) -> dict[str, str]:
    return {
        "X-User-Id": str(external_user_id),
        "X-Account-Id": str(account_id),
    }


def _command(
    client: TestClient,
    external_user_id: int,
    action: str,
    account_id: int = ACCOUNT,
):
    return client.post(
        f"/api/v1/timesheet/{action}",
        headers=_headers(external_user_id, account_id),
        json={"idempotency_key": str(uuid4())},
    )


def _assert_problem(response, status: int, code: str, message: str) -> None:
    assert response.status_code == status, response.text
    error = response.json()["error"]
    assert error["code"] == code
    assert error["message"] == message
    assert any("а" <= char.lower() <= "я" or char.lower() == "ё" for char in message)
    assert UUID(error["request_id"])
    assert response.headers["X-Request-Id"] == error["request_id"]


def _editable_snapshot(body: dict) -> dict:
    return {
        "revision": body["revision"],
        "settings": body["settings"],
        "groups": [
            {key: value for key, value in group.items() if key != "account_id"}
            for group in body["groups"]
        ],
        "users": [
            {
                "amocrm_user_id": user["amocrm_user_id"],
                "track_time": user["track_time"],
                "hide_widget": user["hide_widget"],
                "group_ref": f"id:{user['group_id']}" if user["group_id"] else None,
            }
            for user in body["users"]
        ],
    }


def test_employee_hidden_visible_lifecycle_and_same_day_restart(
    acceptance, monkeypatch
):
    client, db = acceptance
    now = [datetime(2026, 9, 30, 9, 0)]
    monkeypatch.setattr("app.services.timesheet_service.utc_now", lambda: now[0])

    hidden = client.get("/api/v1/timesheet/my-status", headers=_headers(103))
    assert hidden.status_code == 200
    assert (hidden.json()["track_time"], hidden.json()["hide_widget"]) == (True, True)
    db.query(GroupMember).filter_by(account_id=ACCOUNT, user_id=3).one().hide_widget = (
        False
    )
    db.commit()
    assert (
        client.get("/api/v1/timesheet/my-status", headers=_headers(103)).json()[
            "hide_widget"
        ]
        is False
    )

    observed = []
    for at, action, status in [
        (datetime(2026, 9, 30, 9, 0), "start-work", "working"),
        (datetime(2026, 9, 30, 10, 0), "start-break", "on_break"),
        (datetime(2026, 9, 30, 10, 15), "end-break", "working"),
        (datetime(2026, 9, 30, 18, 0), "finish-work", "finished"),
    ]:
        now[0] = at
        response = _command(client, 103, action)
        assert response.status_code == 200, response.text
        assert response.json()["status"] == status
        observed.append(response.json()["session_id"])
    assert len(set(observed)) == 1
    denied = _command(client, 103, "start-work")
    _assert_problem(
        denied,
        409,
        "STATUS_TRANSITION_INVALID",
        "Статус уже изменился. Обновите страницу.",
    )

    group = db.get(WidgetGroup, 10)
    group.allow_restart_session = True
    db.commit()
    now[0] = datetime(2026, 9, 30, 18, 1)
    restarted = _command(client, 103, "start-work")
    assert restarted.status_code == 200, restarted.text
    assert restarted.json()["status"] == "working"
    assert restarted.json()["session_id"] != observed[0]
    assert (
        db.query(WorkSession)
        .filter_by(
            amocrm_account_id=ACCOUNT,
            amocrm_user_id=103,
        )
        .count()
        == 2
    )
    assert (
        db.query(WorkSession)
        .filter(
            WorkSession.amocrm_account_id != ACCOUNT,
            WorkSession.amocrm_user_id == 103,
        )
        .count()
        == 0
    )


def test_manager_is_limited_to_assigned_group_and_employee_detail(acceptance):
    client, _db = acceptance
    headers = _headers(102)
    monitor = client.get("/api/v1/team/status", headers=headers)
    assert monitor.status_code == 200, monitor.text
    assert monitor.json()["viewer"]["role"] == "manager"
    assert {row["id"] for row in monitor.json()["employees"]} == {2, 3}
    assert {row["id"] for row in monitor.json()["groups"]} == {10}

    period = {"from": DAY, "to": DAY}
    own_detail = client.get("/api/v1/team/2/activity", headers=headers, params=period)
    _assert_problem(own_detail, 404, "NOT_FOUND", "Данные не найдены.")
    employee_detail = client.get(
        "/api/v1/team/3/activity", headers=headers, params=period
    )
    assert employee_detail.status_code == 200, employee_detail.text
    assert employee_detail.json()["target"]["id"] == 3
    _assert_problem(
        client.get("/api/v1/team/status?group_id=20", headers=headers),
        404,
        "NOT_FOUND",
        "Данные не найдены.",
    )
    _assert_problem(
        client.get(
            "/api/v1/reports/detailed",
            headers=headers,
            params={"date_from": DAY, "date_to": DAY, "user_id": 4},
        ),
        404,
        "NOT_FOUND",
        "Данные не найдены.",
    )


def test_admin_configuration_monitoring_report_and_safe_excel(acceptance, monkeypatch):
    client, db = acceptance
    monkeypatch.setattr(
        "app.services.timesheet_service.utc_now", lambda: datetime(2026, 9, 30, 9)
    )
    admin_headers = _headers(101)
    read = client.get("/api/v1/settings/snapshot", headers=admin_headers)
    assert read.status_code == 200, read.text
    payload = _editable_snapshot(read.json())
    own_group = next(group for group in payload["groups"] if group["id"] == 10)
    own_group.update(
        work_start_time="08:30:00",
        work_end_time="17:30:00",
        timezone="Europe/Moscow",
        allow_restart_session=True,
    )
    employee = next(user for user in payload["users"] if user["amocrm_user_id"] == 103)
    employee.update(track_time=False, hide_widget=True, group_ref="id:10")
    saved = client.put("/api/v1/settings/snapshot", headers=admin_headers, json=payload)
    assert saved.status_code == 200, saved.text
    saved_group = next(group for group in saved.json()["groups"] if group["id"] == 10)
    assert (
        saved_group["timezone"],
        saved_group["work_start_time"],
        saved_group["allow_restart_session"],
    ) == ("Europe/Moscow", "08:30:00", True)

    saved_employee = next(
        user for user in saved.json()["users"] if user["amocrm_user_id"] == 103
    )
    assert (saved_employee["track_time"], saved_employee["hide_widget"]) == (
        False,
        True,
    )
    disabled_status = client.get("/api/v1/timesheet/my-status", headers=_headers(103))
    assert (
        disabled_status.json()["track_time"],
        disabled_status.json()["hide_widget"],
    ) == (False, True)
    _assert_problem(
        _command(client, 103, "start-work"),
        403,
        "TRACK_TIME_DISABLED",
        "Учёт рабочего времени отключён.",
    )

    enabled_payload = _editable_snapshot(saved.json())
    enabled_employee = next(
        user for user in enabled_payload["users"] if user["amocrm_user_id"] == 103
    )
    enabled_employee.update(track_time=True, hide_widget=False, group_ref="id:10")
    enabled = client.put(
        "/api/v1/settings/snapshot", headers=admin_headers, json=enabled_payload
    )
    assert enabled.status_code == 200, enabled.text
    enabled_employee = next(
        user for user in enabled.json()["users"] if user["amocrm_user_id"] == 103
    )
    assert (enabled_employee["track_time"], enabled_employee["hide_widget"]) == (
        True,
        False,
    )
    enabled_status = client.get("/api/v1/timesheet/my-status", headers=_headers(103))
    assert (
        enabled_status.json()["track_time"],
        enabled_status.json()["hide_widget"],
    ) == (True, False)

    db.get(User, 3).name = '=HYPERLINK("https://evil.invalid")'
    db.commit()
    assert _command(client, 103, "start-work").status_code == 200
    monitor = client.get("/api/v1/team/status", headers=admin_headers)
    assert monitor.status_code == 200, monitor.text
    main_employee = next(row for row in monitor.json()["employees"] if row["id"] == 3)
    assert main_employee["status"] == "working"
    assert all(row["account_id"] == ACCOUNT for row in monitor.json()["employees"])

    report = client.get(
        "/api/v1/reports/detailed",
        headers=admin_headers,
        params={"date_from": DAY, "date_to": DAY, "user_id": 3},
    )
    assert report.status_code == 200, report.text
    assert report.json()["total"] == 1
    assert report.json()["items"][0]["amocrm_user_id"] == 103

    exported = client.post(
        "/api/v1/reports/export-excel",
        headers=admin_headers,
        json={"date_from": DAY, "date_to": DAY, "user_id": 3},
    )
    assert exported.status_code == 200, exported.text
    assert exported.content.startswith(b"PK")
    sheet = load_workbook(BytesIO(exported.content), data_only=False).active
    assert sheet.max_row == 2
    assert sheet["A2"].value == '\'=HYPERLINK("https://evil.invalid")'
    assert (
        db.query(User)
        .filter(User.amocrm_account_id != ACCOUNT, User.amocrm_user_id == 103)
        .count()
        == 40
    )


def test_same_external_user_is_isolated_across_all_40_tenant_apis(
    acceptance, monkeypatch
):
    client, db = acceptance
    monkeypatch.setattr(
        "app.services.timesheet_service.utc_now", lambda: datetime(2026, 9, 30, 9)
    )

    account_ids = list(range(10_000, 10_040))
    for account_id in account_ids:
        status = client.get(
            "/api/v1/timesheet/my-status",
            headers=_headers(103, account_id),
        )
        assert status.status_code == 200, status.text
        assert status.json()["status"] == "not_started"
        assert status.json()["session_id"] is None

        started = _command(client, 103, "start-work", account_id)
        assert started.status_code == 200, started.text
        assert started.json()["status"] == "working"

        current = client.get(
            "/api/v1/timesheet/my-status",
            headers=_headers(103, account_id),
        )
        assert current.status_code == 200, current.text
        assert current.json()["status"] == "working"
        assert current.json()["session_id"] == started.json()["session_id"]

    tenant_sessions = (
        db.query(WorkSession)
        .filter(
            WorkSession.amocrm_account_id.in_(account_ids),
            WorkSession.amocrm_user_id == 103,
        )
        .all()
    )
    assert len(tenant_sessions) == 40
    assert {row.amocrm_account_id for row in tenant_sessions} == set(account_ids)
    assert len({row.id for row in tenant_sessions}) == 40
    assert (
        db.query(WorkSession)
        .filter_by(amocrm_account_id=ACCOUNT, amocrm_user_id=103)
        .count()
        == 0
    )


def test_acceptance_failure_matrix_has_safe_correlated_contracts(
    acceptance, monkeypatch, caplog
):
    from app.core.rate_limit import authenticated_request_limiter
    from app.services.team_service import TeamService

    client, _db = acceptance
    _assert_problem(
        client.get("/api/v1/team/status"),
        401,
        "AMOCRM_TOKEN_EXPIRED",
        "Срок подключения amoCRM истёк.",
    )
    _assert_problem(
        client.get(
            "/api/v1/reports/detailed",
            headers=_headers(103),
            params={"date_from": DAY, "date_to": DAY},
        ),
        403,
        "ACCESS_DENIED",
        "У вас нет доступа к этому разделу.",
    )
    _assert_problem(
        _command(client, 103, "finish-work"),
        409,
        "STATUS_TRANSITION_INVALID",
        "Статус уже изменился. Обновите страницу.",
    )
    _assert_problem(
        client.get(
            "/api/v1/reports/detailed",
            headers=_headers(101),
            params={"date_from": "broken", "date_to": DAY},
        ),
        422,
        "REQUEST_INVALID",
        "Проверьте формат и значения полей запроса.",
    )

    original_rate_check = authenticated_request_limiter.check
    monkeypatch.setattr(
        authenticated_request_limiter, "check", lambda *_args, **_kwargs: (True, 7)
    )
    limited = client.get("/api/v1/team/status", headers=_headers(101))
    _assert_problem(
        limited,
        429,
        "RATE_LIMITED",
        "Слишком много запросов. Повторите попытку позже.",
    )
    assert limited.headers["Retry-After"] == "7"
    monkeypatch.setattr(authenticated_request_limiter, "check", original_rate_check)

    secret = "postgresql://user:never-log@db/private"
    monkeypatch.setattr(
        TeamService,
        "get_monitoring_status",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError(secret)),
    )
    with caplog.at_level(logging.ERROR, logger="app.errors"):
        failed = client.get("/api/v1/team/status", headers=_headers(101))
    _assert_problem(
        failed,
        500,
        "INTERNAL_ERROR",
        "Произошла внутренняя ошибка. Повторите попытку позже.",
    )
    assert secret not in failed.text
    assert "never-log" not in failed.text
    assert secret not in caplog.text
    assert "never-log" not in caplog.text
    assert any(
        getattr(record, "request_id", None) == failed.headers["X-Request-Id"]
        for record in caplog.records
    )

    from app import main

    class UnavailableDatabase:
        def connect(self):
            raise RuntimeError(secret)

    monkeypatch.setattr(main, "database_engine", UnavailableDatabase())
    unavailable = client.get("/health/ready")
    assert unavailable.status_code == 503
    assert UUID(unavailable.headers["X-Request-Id"])
    assert unavailable.json() == {
        "status": "not_ready",
        "checks": {"database": "unavailable", "schema": "unknown"},
    }
    assert secret not in unavailable.text
    assert "never-log" not in unavailable.text
    assert "message" not in unavailable.json()
