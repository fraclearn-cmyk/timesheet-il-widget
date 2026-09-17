"""Behavioral checks for the model/service boundary (PostgreSQL DDL tested separately)."""

from datetime import datetime

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.database import Base
from app.models import User, WorkStatus, CrmEvent
from app.schemas.work_session import WorkSessionResponse
from app.services.session_service import SessionService


@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add_all(
            [
                User(id=7, amocrm_user_id=700, amocrm_account_id=100, name="One"),
            ]
        )
        session.commit()
        yield session
    engine.dispose()


@pytest.mark.filterwarnings("error::DeprecationWarning:app.services.session_service")
@pytest.mark.filterwarnings("error::DeprecationWarning:sqlalchemy.sql.schema")
def test_account_scoped_external_identity_and_session_lifecycle(db):
    db.add(User(id=8, amocrm_user_id=700, amocrm_account_id=200, name="Two"))
    db.commit()
    service = SessionService(db)
    first = service.create_session("100", 700, "<b>One</b>")
    second = service.create_session("200", 700, "Two")
    assert first.amocrm_user_id == 700
    assert first.amocrm_account_id == 100
    assert first.user.id == 7
    assert second.user.id == 8
    assert first.current_status is WorkStatus.WORKING
    assert (
        first.active_duration,
        first.unconfirmed_duration,
        first.break_duration,
        first.idle_duration,
    ) == (0, 0, 0, 0)
    assert service.get_current_session("100", 700).id == first.id
    with pytest.raises(ValueError, match="exists"):
        service.create_session("100", 700, "Duplicate")
    service.update_session(first.id, status="break")
    assert first.current_status is WorkStatus.BREAK
    service.finish_session(first.id)
    assert first.current_status is WorkStatus.FINISHED
    assert first.end_time is not None
    assert service.get_current_session("100", 700) is None
    assert [s.id for s in service.get_user_sessions("100", 700)] == [first.id]
    response = WorkSessionResponse.model_validate(first)
    assert response.user_id == 700  # legacy wire alias is explicitly external
    assert response.active_duration == 0
    with pytest.raises(ValueError, match="ambiguous"):
        service.get_current_session(700)


def test_service_does_not_accept_internal_id_as_external_identity(db):
    with pytest.raises(ValueError, match="user"):
        SessionService(db).create_session("100", 7, "Spoofed")


def test_unknown_crm_author_is_not_a_fake_user(db):
    event = CrmEvent(
        account_id=100,
        external_id="opaque-evt_A1",
        user_id=None,
        event_type="unknown_event",
        occurred_at=datetime(2026, 9, 16),
    )
    db.add(event)
    db.commit()
    db.refresh(event)
    assert event.external_id == "opaque-evt_A1"
    assert event.user_id is None
    assert not event.is_complete


def test_backend_registers_legacy_routes_without_import_failure():
    from app.main import app

    paths = {route.path for route in app.routes}
    assert "/api/v1/sessions/start" in paths
    assert "/api/v1/kpi/my" in paths


@pytest.mark.filterwarnings("error::DeprecationWarning:pydantic.main")
def test_legacy_session_routes_use_external_ids_and_persist_transitions(db):
    from app.api.v1 import sessions
    from app.schemas.work_session import WorkSessionCreate

    created = sessions.start_session(
        WorkSessionCreate(user_id=700, user_name="One"), db
    )
    assert created.user_id == 700
    assert sessions.take_break(700, db).current_status is WorkStatus.BREAK
    assert sessions.resume_work(700, db).current_status is WorkStatus.WORKING
    assert sessions.get_current_session(700, db).id == created.id
    assert sessions.finish_session(700, db).current_status is WorkStatus.FINISHED
    assert sessions.get_session(created.id, db).id == created.id


def test_legacy_consumers_query_canonical_fields_and_correct_account(db):
    from app.models import WorkComment, Department
    from app.services.report_service import ReportService
    from app.services.kpi_service import KPIService
    from app.services.excel_service import ExcelService
    from openpyxl import load_workbook
    from datetime import time

    db.add(
        Department(id=1, name="Sales", work_start_time=time(9), work_end_time=time(18))
    )
    db.get(User, 7).department_id = 1
    db.add(User(id=8, amocrm_user_id=700, amocrm_account_id=200, name="Two"))
    db.commit()
    first = SessionService(db).create_session("100", 700, "One")
    second = SessionService(db).create_session("200", 700, "Two")
    first.total_work_time = 3600
    second.total_work_time = 7200
    db.add(
        WorkComment(
            work_session_id=first.id, author_id=7, author_name="One", comment="Checked"
        )
    )
    db.commit()
    today = first.start_time.date()
    daily = ReportService.get_daily_report(db, "100", today)
    assert len(daily.sessions) == 1
    assert daily.sessions[0].status == "working"
    assert daily.sessions[0].user_id == 700
    assert KPIService(db).calculate_user_kpi(7, 700).hours_today == 1
    assert KPIService(db).calculate_department_kpi(1).hours_today == 1
    output = ExcelService(db).generate_employee_report(7, today, today)
    sheet = load_workbook(output).active
    assert sheet.cell(5, 9).value == "Checked"
    assert sheet.cell(5, 5).value == "1:00"
    department_output = ExcelService(db).generate_department_report(today, today, [1])
    assert load_workbook(department_output).active.cell(5, 9).value == "Checked"


def test_team_reads_normalized_activity_model(db):
    from app.services.team_service import TeamService
    from app.models import CrmEvent

    work = SessionService(db).create_session("100", 700, "One")
    db.add(
        CrmEvent(
            account_id=100,
            external_id="opaque-one",
            user_id=7,
            author_amocrm_user_id=700,
            event_type="lead_added",
            object_type="lead",
            is_complete=1,
            occurred_at=work.start_time,
        )
    )
    db.commit()
    service = TeamService(db)
    assert service.get_team_status()[0]["user_id"] == 700
    assert service.get_team_status_with_rbac(None)[0]["is_online"]
    assert service.get_user_timeline(700)["total_events"] == 1
    assert service.get_user_timeline_history(700)["days"][0]["total_events"] == 1
    assert service.force_finish_session(700, 7, "One", "Done")["success"]


def test_employee_report_event_count_stays_in_account_and_period(db):
    from datetime import timedelta
    from app.models import ActivitySession, ActivityEvent, WorkSession
    from app.services.report_service import ReportService

    db.add(User(id=8, amocrm_user_id=700, amocrm_account_id=200, name="Two"))
    db.commit()
    first = SessionService(db).create_session("100", 700, "One")
    second = SessionService(db).create_session("200", 700, "Two")
    old = WorkSession(
        amocrm_account_id=100,
        amocrm_user_id=700,
        user_name="One",
        start_time=first.start_time - timedelta(days=2),
    )
    db.add(old)
    db.flush()
    for work in (first, second, old):
        activity = ActivitySession(
            work_session_id=work.id,
            entity_type="lead",
            entity_id=1,
            start_time=work.start_time,
        )
        activity.activity_events.append(ActivityEvent(event_type="card_updated"))
        db.add(activity)
    db.commit()
    today = first.start_time.date()
    report = ReportService.get_employee_report(db, "100", 700, today, today)
    assert report.activities[0].event_count == 1
