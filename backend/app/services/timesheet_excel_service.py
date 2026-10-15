"""Allowlisted XLSX rendering for detailed timesheet rows."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.schemas.report import DetailedReportRow, TimesheetColumn


HEADERS = {
    TimesheetColumn.EMPLOYEE: "Сотрудник",
    TimesheetColumn.DATE: "Дата",
    TimesheetColumn.START: "Начало",
    TimesheetColumn.END: "Окончание",
    TimesheetColumn.BREAK: "Перерыв",
    TimesheetColumn.WORK: "Работа",
    TimesheetColumn.SHIFT_LATENESS: "Минуты",
    TimesheetColumn.SHIFT_LATE_REASON: "Объяснительная",
    TimesheetColumn.BREAK_LATENESS: "Минуты",
    TimesheetColumn.BREAK_LATE_REASON: "Объяснительная",
    TimesheetColumn.STATUS: "Статус",
}
STATUS_LABELS = {
    "working": "Работает",
    "on_break": "Перерыв",
    "finished": "Закончил(а)",
}
DATETIME_FORMAT = "dd.mm.yyyy hh:mm"
DATE_FORMAT = "dd.mm.yyyy"
DURATION_FORMAT = "[h]:mm"
LATENESS_GROUPS = {
    TimesheetColumn.SHIFT_LATENESS: "Начало дня",
    TimesheetColumn.SHIFT_LATE_REASON: "Начало дня",
    TimesheetColumn.BREAK_LATENESS: "Перерыв",
    TimesheetColumn.BREAK_LATE_REASON: "Перерыв",
}
REASON_COLUMNS = {TimesheetColumn.SHIFT_LATE_REASON, TimesheetColumn.BREAK_LATE_REASON}
MINUTE_COLUMNS = {TimesheetColumn.SHIFT_LATENESS, TimesheetColumn.BREAK_LATENESS}


def _header(sheet, columns: list[TimesheetColumn]) -> int:
    if not any(column in LATENESS_GROUPS for column in columns):
        sheet.append([HEADERS[column] for column in columns])
        return 1
    for index, column in enumerate(columns, 1):
        if column not in LATENESS_GROUPS:
            sheet.cell(1, index, HEADERS[column])
            sheet.merge_cells(
                start_row=1, end_row=3, start_column=index, end_column=index
            )
        else:
            sheet.cell(3, index, HEADERS[column])
    # Arbitrary selections keep their requested order; merge only contiguous groups.
    for level in (1, 2):
        start = 1
        while start <= len(columns):
            column = columns[start - 1]
            if column not in LATENESS_GROUPS:
                start += 1
                continue
            label = "Опоздания" if level == 1 else LATENESS_GROUPS[column]
            end = start
            while (
                end < len(columns)
                and columns[end] in LATENESS_GROUPS
                and (level == 1 or LATENESS_GROUPS[columns[end]] == label)
            ):
                end += 1
            sheet.cell(level, start, label)
            if end > start:
                sheet.merge_cells(
                    start_row=level, end_row=level, start_column=start, end_column=end
                )
            start = end + 1
    return 3


def _break_reasons(row: DetailedReportRow) -> str:
    lines = []
    for incident in row.break_late_reason:
        label = (
            f"Перерыв №{incident.sequence} — {(incident.late_seconds + 59) // 60} мин"
        )
        lines.append(f"{label}: {incident.reason}" if incident.reason else label)
    return _safe_text("\n".join(lines))


def safe_report_filename(date_from: date, date_to: date) -> str:
    return f"timesheet_{date_from.isoformat()}_{date_to.isoformat()}.xlsx"


def _safe_text(value: str) -> str:
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _local_datetime(value: datetime | None, group_timezone: str) -> datetime | None:
    if value is None:
        return None
    utc_value = (
        value.replace(tzinfo=timezone.utc)
        if value.tzinfo is None
        else value.astimezone(timezone.utc)
    )
    return utc_value.astimezone(ZoneInfo(group_timezone)).replace(tzinfo=None)


class TimesheetExcelService:
    def render(
        self, rows: list[DetailedReportRow], columns: list[TimesheetColumn]
    ) -> bytes:
        book = Workbook()
        sheet = book.active
        sheet.title = "Табель"
        header_rows = _header(sheet, columns)
        for row in rows:
            values = {
                TimesheetColumn.EMPLOYEE: _safe_text(row.employee_name),
                TimesheetColumn.DATE: row.date,
                TimesheetColumn.START: _local_datetime(
                    row.started_at, row.group_timezone
                ),
                TimesheetColumn.END: _local_datetime(row.ended_at, row.group_timezone),
                TimesheetColumn.BREAK: row.break_seconds / 86400,
                TimesheetColumn.WORK: row.work_seconds / 86400,
                TimesheetColumn.SHIFT_LATENESS: (row.shift_late_seconds + 59) // 60,
                TimesheetColumn.SHIFT_LATE_REASON: _safe_text(
                    row.shift_late_reason or ""
                ),
                TimesheetColumn.BREAK_LATENESS: (row.break_late_seconds + 59) // 60,
                TimesheetColumn.BREAK_LATE_REASON: _break_reasons(row),
                TimesheetColumn.STATUS: STATUS_LABELS[row.status],
            }
            sheet.append([values[column] for column in columns])
        for cells in sheet.iter_rows(max_row=header_rows):
            for cell in cells:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.font = Font(bold=True)
        for column_index, column in enumerate(columns, start=1):
            number_format = None
            if column == TimesheetColumn.DATE:
                number_format = DATE_FORMAT
            elif column in (TimesheetColumn.START, TimesheetColumn.END):
                number_format = DATETIME_FORMAT
            elif column in (
                TimesheetColumn.BREAK,
                TimesheetColumn.WORK,
            ):
                number_format = DURATION_FORMAT
            elif column in MINUTE_COLUMNS:
                number_format = "0"
            if number_format is not None:
                for cells in sheet.iter_cols(
                    min_col=column_index, max_col=column_index, min_row=header_rows + 1
                ):
                    for cell in cells:
                        cell.number_format = number_format
            if column in REASON_COLUMNS:
                for cells in sheet.iter_cols(
                    min_col=column_index, max_col=column_index, min_row=header_rows + 1
                ):
                    for cell in cells:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions[get_column_letter(column_index)].width = (
                45 if column in REASON_COLUMNS else 20
            )
        sheet.freeze_panes = f"A{header_rows + 1}"
        sheet.auto_filter.ref = (
            f"A{header_rows}:{get_column_letter(len(columns))}{sheet.max_row}"
        )
        output = BytesIO()
        book.save(output)
        return output.getvalue()
